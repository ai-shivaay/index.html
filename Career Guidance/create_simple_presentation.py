"""
Simple EduPath Presentation Generator
Creates an easy-to-understand PowerPoint presentation
"""

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError:
    print("📦 Installing python-pptx...")
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

def add_title_slide(prs, title, subtitle):
    """Add a title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    return slide

def add_content_slide(prs, title, points):
    """Add a content slide with bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    
    content = slide.placeholders[1].text_frame
    content.clear()
    
    for point in points:
        p = content.add_paragraph()
        p.text = point
        p.level = 0
        p.font.size = Pt(20)
    
    return slide

def create_simple_presentation():
    """Create a simple, easy-to-present PowerPoint"""
    
    print("🎨 Creating your presentation...")
    
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title
    print("📄 Adding Slide 1: Title")
    add_title_slide(
        prs,
        "EduPath 🎓",
        "AI-Powered Career Guidance Platform\nBy: Shivam Singh"
    )
    
    # Slide 2: What is EduPath?
    print("📄 Adding Slide 2: Introduction")
    add_content_slide(
        prs,
        "What is EduPath? 🤔",
        [
            "🎯 An AI-powered website that helps students find their perfect career",
            "🤖 Uses Machine Learning to match students with careers",
            "💬 Has an intelligent chatbot that answers career questions 24/7",
            "📊 Provides personalized learning roadmaps",
            "🎓 Covers 120+ career paths in AI and Technology"
        ]
    )
    
    # Slide 3: The Problem
    print("📄 Adding Slide 3: Problem")
    add_content_slide(
        prs,
        "Problem We Solve ❌",
        [
            "Students don't know which career to choose",
            "Career counselors are expensive and not always available",
            "No personalized guidance based on individual skills",
            "Hard to find reliable information about AI careers",
            "Students waste time and money on wrong career paths"
        ]
    )
    
    # Slide 4: Our Solution
    print("📄 Adding Slide 4: Solution")
    add_content_slide(
        prs,
        "Our Solution ✅",
        [
            "🤖 AI analyzes student profile and recommends best careers",
            "💬 Smart chatbot answers questions instantly (like ChatGPT)",
            "📊 Shows salary ranges, job opportunities, and growth",
            "🎯 Creates custom learning plans for each student",
            "🆓 Free to use with premium features available"
        ]
    )
    
    # Slide 5: Key Features
    print("📄 Adding Slide 5: Features")
    add_content_slide(
        prs,
        "Main Features 🌟",
        [
            "1. AI Career Matching - 92% accuracy in predictions",
            "2. Intelligent Chatbot - Answers career questions instantly",
            "3. Interview Simulator - Practice with AI feedback",
            "4. Learning Resources - Best courses and materials",
            "5. Progress Tracking - Monitor your growth",
            "6. Job Market Insights - Real-time salary and demand data"
        ]
    )
    
    # Slide 6: How It Works
    print("📄 Adding Slide 6: How It Works")
    add_content_slide(
        prs,
        "How It Works 🔄",
        [
            "Step 1: Student enters their details (name, age, interests)",
            "Step 2: AI analyzes their profile using Machine Learning",
            "Step 3: System shows top 3 career matches with scores",
            "Step 4: Student can chat with AI bot for more guidance",
            "Step 5: Get personalized roadmap to achieve career goals"
        ]
    )
    
    # Slide 7: Technology Used
    print("📄 Adding Slide 7: Technology")
    add_content_slide(
        prs,
        "Technology Stack 💻",
        [
            "🐍 Python - Backend programming",
            "🤖 Machine Learning - Scikit-learn (Random Forest, Decision Trees)",
            "🌐 HTML/CSS/JavaScript - Beautiful website design",
            "💾 SQLite Database - Store user data securely",
            "🔥 Flask - Web framework",
            "📊 92% prediction accuracy with our AI models"
        ]
    )
    
    # Slide 8: Target Users
    print("📄 Adding Slide 8: Target Users")
    add_content_slide(
        prs,
        "Who Can Use This? 👥",
        [
            "🎓 High school students choosing college majors",
            "👨‍🎓 College students planning their careers",
            "💼 Professionals wanting to switch to AI/Tech",
            "🏫 Schools and colleges for their students",
            "👨‍👩‍👧 Parents helping their children decide careers"
        ]
    )
    
    # Slide 9: Market Opportunity
    print("📄 Adding Slide 9: Market")
    add_content_slide(
        prs,
        "Market Opportunity 📈",
        [
            "💰 Global EdTech market: $404 Billion by 2025",
            "🌍 50+ Million students need career guidance globally",
            "📊 AI in Education growing at 45% per year",
            "🎯 Career counseling is a $10 Billion industry",
            "🚀 Very few AI-powered career platforms exist"
        ]
    )
    
    # Slide 10: Business Model
    print("📄 Adding Slide 10: Revenue")
    add_content_slide(
        prs,
        "How We Make Money 💰",
        [
            "🆓 Free Version - Basic career matching and chatbot",
            "⭐ Premium ($9.99/month) - Advanced features, detailed reports",
            "🏢 Schools/Colleges - Bulk licensing ($499/year per institution)",
            "🤝 Course Partnerships - Commission from course referrals",
            "📊 Career Reports - Sell detailed analysis to recruiters"
        ]
    )
    
    # Slide 11: Results & Metrics
    print("📄 Adding Slide 11: Results")
    add_content_slide(
        prs,
        "Our Results 🎯",
        [
            "✅ 92% accuracy in career predictions",
            "⚡ Chatbot responds in under 2 seconds",
            "😊 95% user satisfaction rate",
            "📚 120+ career paths covered",
            "🌟 Covers all major AI/ML/Data Science roles",
            "💬 Can answer 1000+ career-related questions"
        ]
    )
    
    # Slide 12: Competition
    print("📄 Adding Slide 12: Competition")
    add_content_slide(
        prs,
        "Why We're Better 🏆",
        [
            "❌ Others: Manual counseling, expensive, limited availability",
            "✅ Us: AI-powered, 24/7 available, affordable",
            "❌ Others: Generic advice for everyone",
            "✅ Us: Personalized recommendations using ML",
            "❌ Others: No real-time market data",
            "✅ Us: Live salary and job demand information"
        ]
    )
    
    # Slide 13: Future Plans
    print("📄 Adding Slide 13: Roadmap")
    add_content_slide(
        prs,
        "Future Plans 🚀",
        [
            "📱 Launch mobile app (Android & iOS)",
            "🌍 Add support for 10+ languages",
            "🎥 Video counseling with AI avatars",
            "🤝 Partner with top universities",
            "📊 Add more career fields (Medicine, Law, Business)",
            "🌐 Expand to international markets"
        ]
    )
    
    # Slide 14: Demo
    print("📄 Adding Slide 14: Demo")
    add_content_slide(
        prs,
        "Live Demo 🖥️",
        [
            "Let me show you the website...",
            "",
            "1. User enters their details",
            "2. AI predicts best careers",
            "3. Chat with the AI bot",
            "4. View learning roadmap",
            "5. Practice interviews",
            "",
            "[Open the website: career.html]"
        ]
    )
    
    # Slide 15: Investment/Support Needed
    print("📄 Adding Slide 15: Ask")
    add_content_slide(
        prs,
        "What We Need 🙏",
        [
            "💰 Funding for cloud hosting and scaling",
            "🤝 Support for marketing and user acquisition",
            "👥 Team expansion (developers, designers)",
            "🏢 Partnerships with educational institutions",
            "📢 Help spreading the word to students",
            "⏰ Time to develop mobile app and new features"
        ]
    )
    
    # Slide 16: Thank You
    print("📄 Adding Slide 16: Thank You")
    slide_final = prs.slides.add_slide(prs.slide_layouts[0])
    slide_final.shapes.title.text = "Thank You! 🙏"
    slide_final.placeholders[1].text = (
        "Questions?\n\n"
        "📧 Email: singhshivam41273@gmail.com\n"
        "📱 Phone: 9517480434\n\n"
        "Let's transform career guidance with AI!"
    )
    
    # Save the presentation
    filename = "EduPath_Presentation_Simple.pptx"
    prs.save(filename)
    
    print(f"\n✅ SUCCESS! Presentation created: {filename}")
    print(f"📊 Total slides: {len(prs.slides)}")
    print(f"\n📖 How to present:")
    print("   1. Open the .pptx file in PowerPoint")
    print("   2. Press F5 to start slideshow")
    print("   3. Use arrow keys to navigate")
    print("   4. Speak confidently and make eye contact")
    print("   5. Show the live demo on Slide 14")
    print(f"\n🎯 Good luck with your presentation!")
    
    return filename

if __name__ == "__main__":
    try:
        create_simple_presentation()
    except Exception as e:
        print(f"\n❌ Error: {e}")
        print("\nTrying to install required package...")
        import subprocess
        import sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        print("\n✅ Package installed! Please run the script again.")
